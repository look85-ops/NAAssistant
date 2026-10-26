from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path('wardrobe/2026-10-12_капсула-тихий-архитектор.xlsx')
blue = PatternFill('solid', fgColor='4472C4')
green = PatternFill('solid', fgColor='C6E0B4')
white = Font(color='FFFFFF', bold=True)
thin = Side(style='thin', color='A6A6A6')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

items = [
 ['Mango','Жакет BERGEN','жакет','серый',12999,'63% полиэстер, 32% вискоза, 5% эластан','костюмный текстиль; архитектурный слой','44/46 RUS (M), 18 шт.','https://www.lamoda.ru/p/rtlafk737101/clothes-mango-zhaket/'],
 ["Looker’s",'Брюки Гонконг графит','брюки','графит',8000,'100% лен','широкие; стандартная посадка; 102 см (XS)','44/46 RUS (M), 9 шт.','https://www.lamoda.ru/p/xd001xw05knf/clothes-lookers-bryuki/'],
 ['Jack Wolfskin','NATURE CORDUROY PANTS W','брюки','черный',13319,'100% органический хлопок','вельвет; утилитарный плотный низ','44 RUS — 2 шт.; 46 RUS — 2 шт.','https://www.lamoda.ru/p/rtlafc627901/clothes-jackwolfskin-bryuki/'],
 ['2Mood','Свободная рубашка','рубашка','белый/полоска',9990,'58% хлопок, 37% полиамид, 5% спандекс','свободный силуэт; спущенное плечо','44/46 RUS (M), 11 шт.','https://www.lamoda.ru/p/mp002xw1dhw6/clothes-2mood-rubashka/'],
 ['Luisa Wang','Кардиган','трикотаж','ржаво-коричневый',10990,'100% шерсть','тонкий ровный трикотаж; прямой мягкий силуэт','44 RUS — 3 шт.; 46 RUS — 1 шт.','https://www.lamoda.ru/p/mp002xw1k4c1/clothes-luisawang-kardigan/'],
 ["Levi’s®",'Свитшот','свитшот','серый',10619,'80% хлопок, 20% полиэстер','толстовочный трикотаж; съемочный слой','44/46 RUS (M), 8 шт.','https://www.lamoda.ru/p/rtlafb808801/clothes-levis-svitshot/'],
 ['Black Pine','Платье-поло миди','платье','голубой',29410,'100% мериносовая шерсть extrafine','прямой крой; поло; длина 117 см','OneSize 42/46 RUS, 4 шт.','https://www.lamoda.ru/p/mp002xw1e1xg/clothes-blackpine-plate/'],
 ['Mango','Футболка JENNI','футболка','мультиколор/полоска',2799,'59% полиэстер, 25% лиоцелл, 10% полиамид, 3% шерсть, 3% эластан','тонкий связующий слой','44/46 RUS (M), 20 шт.','https://www.lamoda.ru/p/rtlafo839701/clothes-mango-futbolka/'],
]

looks = [
 ['Собеседование без униформы','жакет + рубашка + графитовые брюки','интервью; офис'],
 ['Продуктовый руководитель','жакет + JENNI + черный вельвет','интервью; офис; кафе'],
 ['Монтажный день','свитшот + черный вельвет + JENNI','БГАИ/съемки; прогулки'],
 ['Разбор проекта','рубашка навыпуск + черный вельвет + кардиган','офис; БГАИ'],
 ['Тихая премьера','голубое платье + жакет','театр; галерея; мероприятие'],
 ['Галерейный день','голубое платье + кардиган','галерея; кафе; прогулки'],
 ['Городская редактура','кардиган + JENNI + графитовые брюки','офис; кафе; прогулки'],
 ['Съемка → встреча','рубашка под свитшотом + графитовые брюки','БГАИ/съемки; офис'],
 ['Спокойная пятница','кардиган + рубашка + черный вельвет','офис; театр; прогулки'],
]

checks = [
 [r[0]+' '+r[1],r[7],'Подтверждено карточкой 12.10.2026',('Проверить на просвет' if 'Looker' in r[0] else 'Проверить посадку, длину и свободу'),r[8]] for r in items
]

anti = [
 ['Тотальный черный','слишком буквальный образ режиссера'],['Бежевый костюм','бесхарактерный минимализм'],
 ['Винный/изумрудный акцент','не переносить прежнюю цветовую логику'],['Lime кружевной жакет','приталенный силуэт «песочные часы»'],
 ['DIVNO жилет','нет 44/46; размеры с 50/52'],['12 STOREEZ Flint','44 и 46 распроданы'],
 ['Pennyblack PBBBOSCO','44/46 распродан'],['Обувь и сумка','нет данных о размере обуви; система уже полна'],
 ['Пальто, джинсы, белье','исключено условиями задачи'],
]

def add_sheet(wb, name, title, headers, rows):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws['A1'].font = Font(bold=True,size=14,color='FFFFFF'); ws['A1'].fill = blue
    ws.append(headers)
    for c in ws[2]: c.fill=blue; c.font=white; c.border=border; c.alignment=Alignment(wrap_text=True)
    for row in rows:
        ws.append(row)
        for c in ws[ws.max_row]: c.border=border; c.alignment=Alignment(wrap_text=True,vertical='top')
        for c in ws[ws.max_row]:
            if isinstance(c.value,str) and c.value.startswith('http'):
                c.hyperlink=c.value; c.style='Hyperlink'
    ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{get_column_letter(len(headers))}{ws.max_row}'
    for col in range(1,len(headers)+1):
        width=max(len(str(ws.cell(r,col).value or '')) for r in range(1,ws.max_row+1))+2
        ws.column_dimensions[get_column_letter(col)].width=min(max(width,12),45)
    return ws

wb=Workbook(); del wb['Sheet']
ws=add_sheet(wb,'Капсула','Тихий архитектор — 98 126 ₽',['Бренд','Вещь','Категория','Цвет','Цена, ₽','Состав','Крой/плотность','Наличие размера','Ссылка'],items)
ws.append(['ИТОГО','','','',sum(x[4] for x in items),'','','',''])
for c in ws[ws.max_row]: c.fill=green; c.font=Font(bold=True); c.border=border
add_sheet(wb,'Луки','9 переходных луков',['Название','Состав','Контексты'],looks)
add_sheet(wb,'Проверка размеров','Размеры проверены в карточках Lamoda',['Позиция','Фактическое наличие','Статус','Что проверить на примерке','Ссылка'],checks)
add_sheet(wb,'Анти-список','Что намеренно не выбрано',['Не выбрано','Почему'],anti)
wb.save(OUT)
print(OUT)
