# Собирает .md -> .docx «Обязанности руководителей ЭЧ» с выделением «общего» (жёлтая заливка).
# Запуск: C:\Users\marcenuk\AppData\Local\Python\pythoncore-3.14-64\python.exe scripts/build_duties_docx.py <input.md> <output.docx>

import sys, re
sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
HIGHLIGHT = 'FFEB9C'


def shade(paragraph, fill=HIGHLIGHT):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def run(paragraph, text, bold=False, italic=False, size=11, color=None):
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def add_bullet(doc, text, common=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    if common:
        shade(p)
        run(p, 'ОБЩЕЕ: ', bold=True, size=10)
    run(p, text, size=10)
    return p


def convert(md_path, docx_path):
    doc = Document()
    title_done = False

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue

        if line.startswith('# ') and not title_done:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            run(p, line[2:].strip(), bold=True, size=16, color=ACCENT)
            title_done = True
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run(p, line[3:].strip(), bold=True, size=13, color=ACCENT)
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run(p, line[4:].strip(), bold=True, size=11)
            continue

        if line.startswith('---'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run('─' * 40)
            continue

        if line.startswith('+ ОБЩЕЕ: '):
            add_bullet(doc, line[len('+ ОБЩЕЕ: '):].lstrip(), common=True)
            continue

        if line.startswith('- '):
            add_bullet(doc, line[2:].strip())
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run(p, line, size=10)

    doc.save(docx_path)
    print(f"OK: {docx_path}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python build_duties_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])