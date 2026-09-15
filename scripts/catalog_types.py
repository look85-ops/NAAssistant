# -*- coding: utf-8 -*-
"""
Catalog slide templates in a presentation.
Usage: python catalog_types.py <input.pptx> <output.json>
"""
import json
import sys
from pptx import Presentation
from pptx.util import Emu

if len(sys.argv) >= 3:
    CUR = sys.argv[1]
    OUT = sys.argv[2]
else:
    CUR = r"C:\Users\marcenuk\Desktop\П+Путь_М1_Организация хозяйства пути_26\М1_Д1_26\Управляющая_презентация_М1Д1_26.pptx"
    OUT = r"C:\Users\marcenuk\AppData\Local\Temp\opencode\catalog.json"

prs = Presentation(CUR)
sw = prs.slide_width
sh = prs.slide_height
print(f"slide size: {sw} x {sh} EMU  ({sw/914400:.2f} x {sh/914400:.2f} inch)")

def first_line(slide):
    for shp in slide.shapes:
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    return t
    return ""

def all_text(slide):
    out = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    out.append(t)
    return out

def classify(slide, layout_name, texts):
    joined = " | ".join(texts).lower()
    if not texts:
        return "EMPTY"
    if layout_name == "fbv тема":
        return "SECTION_HEADER"
    if "перерыв" in joined and "мин" in joined:
        return "BREAK"
    if "самостоятельная работа" in joined:
        return "SELFWORK"
    if ("модуль" in joined or "модуль" in texts[0].lower() if texts else False) and ("обсудим" in joined or "поразмышляем" in joined):
        return "DISCUSS"
    if layout_name == "Фирменная рамка №1 (макс)":
        if "?" in joined:
            # determine if question or discussion
            qmarks = joined.count("?")
            if qmarks >= 3:
                return "DISCUSS"
            return "QUESTION_MOD1"
        return "FRAMED"
    if layout_name == "Пустой слайд":
        return "COVER" if "профессиональн" in joined or "развитие" in joined else "BLANK"
    if "фирменная рамка №2" == layout_name.lower():
        return "SECTION_INTRO"
    return "CONTENT"

catalog = []
for i, s in enumerate(prs.slides, 1):
    layout = s.slide_layout.name if s.slide_layout else "?"
    texts = all_text(s)
    fl = first_line(s)
    typ = classify(s, layout, texts)
    catalog.append({
        "n": i,
        "layout": layout,
        "type": typ,
        "title": fl[:80],
        "texts": texts[:10],
    })

with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"slide_size_emu": [sw, sh], "slides": catalog}, f, ensure_ascii=False, indent=1)
print("saved", OUT)

from collections import Counter, defaultdict
counter = Counter(c["type"] for c in catalog)
firsts = defaultdict(list)
for c in catalog:
    if len(firsts[c["type"]]) < 3:
        firsts[c["type"]].append(c["n"])

print("--- types ---")
for typ, cnt in counter.most_common():
    print(f"  {typ:16s} n={cnt:3d}  examples: {firsts[typ]}")
