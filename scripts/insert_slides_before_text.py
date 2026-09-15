"""Insert selected slide PNGs immediately before matched DOCX paragraphs."""
import argparse
import json
import os
import zipfile
from docx import Document
from docx.shared import Cm

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('slide_dir')
    ap.add_argument('mapping')
    ap.add_argument('--width-cm', default='8.8')
    args = ap.parse_args()

    with open(args.mapping, encoding='utf-8') as f:
        mapping = json.load(f)['paragraph_to_slide']

    doc = Document(args.docx)
    paragraphs = list(doc.element.body.iter('{%s}p' % W))
    inserted = 0
    for p_index, slide_no in sorted(mapping.items(), key=lambda x: int(x[0]), reverse=True):
        p = paragraphs[int(p_index)]
        png = os.path.join(args.slide_dir, 'slide_%03d.png' % int(slide_no))
        if not os.path.exists(png):
            continue
        np = doc.add_paragraph()
        np.alignment = 1
        np.add_run().add_picture(png, width=Cm(float(args.width_cm)))
        p.addprevious(np._p)
        inserted += 1

    doc.save(args.docx)
    print('inserted', inserted)


if __name__ == '__main__':
    main()
