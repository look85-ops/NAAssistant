# -*- coding: utf-8 -*-
"""Фиксер опечаток для docx-сценариев П+П.

Запуск:
    py fix_typos.py "путь\к\файлу.docx"

Сохраняет исправленную копию рядом с оригиналом как "<имя>_fixed.docx".
Оригинал не изменяется.
"""
import sys, io, os
import docx

# (подстрока-искомая, подстрока-замена) — порядок важен: сначала длинные/частные
REPLACEMENTS = [
    ("инновационного подход к", "инновационного подхода к"),
    ("земляного полотна земляного полотна", "земляного полотна"),
    ("с уставленными скоростями", "с установленными скоростями"),
    ("для вашего предпритятия", "для вашего предприятия"),
    ("Раздел торой", "Раздел второй"),
    ("остродефеткн", "остродефектн"),
    ("прядка", "порядка"),
    ("совей", "своей"),
]


def fix_paragraph(p):
    """Заменяет опечатки в параграфе. Возвращает число замен."""
    count = 0
    for old, new in REPLACEMENTS:
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                count += 1
    return count


def fix_full_text(p):
    """Fallback: если слово разбито по runs — правим весь текст параграфа."""
    text = p.text
    fixed = text
    for old, new in REPLACEMENTS:
        fixed = fixed.replace(old, new)
    if fixed != text and p.runs:
        p.runs[0].text = fixed
        for r in p.runs[1:]:
            r.text = ""
        return 1
    return 0


def process(doc):
    total = 0
    # параграфы документа
    for p in doc.paragraphs:
        total += fix_paragraph(p) + fix_full_text(p)
    # параграфы внутри таблиц
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    total += fix_paragraph(p) + fix_full_text(p)
    return total


def main():
    if len(sys.argv) < 2:
        print("Укажи путь к docx: py fix_typos.py <file.docx>")
        sys.exit(1)
    src = sys.argv[1]
    if not os.path.exists(src):
        print("Файл не найден:", src)
        sys.exit(1)
    base, ext = os.path.splitext(src)
    dst = base + "_fixed" + ext

    doc = docx.Document(src)
    n = process(doc)
    doc.save(dst)
    print("Замен:", n)
    print("Сохранено:", dst)


if __name__ == "__main__":
    main()
