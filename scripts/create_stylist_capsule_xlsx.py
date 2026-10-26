import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
head_fill = PatternFill('solid', fgColor='5B4A3F')
head_font = Font(bold=True, color='FFFFFF')
border = Border(*(Side(style='thin'),) * 4)

def sheet(name, title, headers, rows):
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for col, value in enumerate(headers, 1):
        cell = ws.cell(3, col, value)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True)
        cell.border = border
    for row_no, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row_no, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = border
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 38
    return ws

del wb['Sheet']
sheet('Капсула', 'Самостоятельная капсула: 88 794 ₽', ['Вещь', 'Цвет', 'Цена, ₽', 'Ссылка'], [
    ['DIVNO жилет «План Б»', 'бордовый', 13990, 'https://www.lamoda.ru/p/mp002xw1p3mf/clothes-divno-zhilet/'],
    ['Arny Praht TIPSY 2.0', 'зелёный', 7482, 'https://www.lamoda.ru/p/mp002xw1ehxl/bags-arnypraht-sumka/'],
    ['12 STOREEZ лонгслив Flint', 'синий', 7800, 'https://www.lamoda.ru/p/mp002xw1p17p/clothes-12storeez-longsliv/'],
    ['Jack Wolfskin брюки NATURE CORDUROY PANTS W', 'чёрный', 11988, 'https://www.lamoda.ru/p/rtlafc627901/clothes-jackwolfskin-bryuki/'],
    ['Luisa Wang кардиган', 'шоколад', 10221, 'https://www.lamoda.ru/p/mp002xw1k4c1/clothes-luisawang-kardigan/'],
    ['Mango футболка JENNI', 'серо-синяя полоска', 2380, 'https://www.lamoda.ru/p/rtlafo839701/clothes-mango-futbolka/'],
    ['Pennyblack джемпер PBBBOSCO', 'серый', 16915, 'https://www.lamoda.ru/p/rtlafl423101/clothes-pennyblack-dzhemper/'],
    ["Levi's свитшот", 'серый', 9027, 'https://www.lamoda.ru/p/rtlafb808801/clothes-levis-svitshot/'],
    ['2Mood рубашка', 'белый', 8991, 'https://www.lamoda.ru/p/mp002xw1dhw6/clothes-2mood-rubashka/'],
])
sheet('Луки', 'Луки капсулы', ['Название', 'Состав', 'Роль'], [
    ['Театральный куратор', 'белая рубашка + бордовый жилет + чёрный вельвет', 'цвет и вертикаль'],
    ['Библиотека', 'синий лонгслив + бордовый жилет + шоколадный кардиган', 'интеллектуальный цвет'],
    ['Мягкая сила', 'серый джемпер + бордовый жилет + зелёная сумка', 'статусный акцент'],
    ['Студия', 'серый свитшот + вельвет + зелёная сумка', 'расслабленная графика'],
    ['Полоска как фон', 'Mango под белую рубашку + бордовый жилет', 'слой и рисунок'],
])
path = os.path.join('wardrobe', '2026-10-12_самостоятельная-капсула-стилиста.xlsx')
wb.save(path)
print(path)
